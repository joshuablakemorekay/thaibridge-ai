"""Build the native speaker's Pronunciation Analysis test sheet.

The tone feedback cannot be checked by reading or listening — only by using
it. So this sheet is a short set of trials: say each word RIGHT, then say it
with a WRONG tone on purpose, and write down what 🎤 Say it reported. A fair
checker gives the right tone a high score and the wrong one a low score with
advice that names the real difference.

The words are one-syllable words from the vocabulary topics, grouped by their
tone, so each of the five Thai tones is tried several times. Mid against low
is listed on purpose: the checker is known to find those two hard to tell
apart (they differ mostly in pitch height, which it sets aside to treat every
voice fairly), and a native speaker trying it is the best measure of how much
that matters.

    python scripts/make_pronunciation_review.py
"""

import argparse
import os
import sys
import unicodedata

os.environ["DATABASE_URL"] = ""

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, REPO_ROOT)

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

import app  # noqa: E402
import thai_audio  # noqa: E402

THAI_FONT = "Leelawadee UI"
PER_TONE = 4      # words tried per tone
TONE_TH = {"mid": "สามัญ", "low": "เอก", "falling": "โท", "high": "ตรี", "rising": "จัตวา"}
TONE_ORDER = ["mid", "low", "falling", "high", "rising"]

INTRO_TH = [
    "สวัสดีครับ 🙏",
    "ชุดนี้เป็นการทดลองใช้ \"🎤 Say it\" (วิเคราะห์การออกเสียง) ในหน้าคำศัพท์ของแอป",
    "วิธีทดลอง: เปิดหน้าคำศัพท์ที่บอกในช่อง หน้า กด 🎤 Say it แล้วกด Next word จนเจอคำนั้น "
    "1) พูดคำให้ถูก จดคะแนนที่ได้ 2) พูดผิดวรรณยุกต์โดยตั้งใจ (เช่น ช่อง พูดผิดเป็น) จดคะแนนอีกครั้ง",
    "ช่อง ยุติธรรมไหม ให้เลือก ✓ ถ้าพูดถูกได้คะแนนสูงและพูดผิดได้คะแนนต่ำ ✗ ถ้าไม่ใช่ "
    "และเขียนในช่อง หมายเหตุ ถ้าคำแนะนำไม่ตรงกับที่พูด หรือระบบไม่ได้ยินคำที่ถูก",
    "ขอบพระคุณมากครับ",
]
INTRO_EN = [
    "This sheet is a hands-on test of 🎤 Say it (Pronunciation Analysis) on the vocabulary pages.",
    "How: open the page in the Page column, tap 🎤 Say it, and tap Next word until that word "
    "comes up. (1) Say it correctly and write down the score. (2) Say it with the wrong tone "
    "on purpose (the 'Say it wrongly as' column) and write down that score.",
    "Under Fair?, choose ✓ if the right tone scored high and the wrong tone scored low, or ✗ if "
    "not. Use the Note column if the advice didn't match what you said, or the word check "
    "didn't recognise the correct word.",
    "Mid and low tones are known to be hard for the checker to tell apart, so those rows "
    "matter most.",
]

HEADERS = ["#", "หน้า · Page", "ภาษาไทย · Thai", "ThaiBridge Spelling", "วรรณยุกต์ · Tone",
           "พูดผิดเป็น · Say it wrongly as", "คะแนนตอนพูดถูก · Score said right",
           "คะแนนตอนพูดผิด · Score said wrong", "ยุติธรรมไหม · Fair?", "หมายเหตุ · Note"]
WIDTHS = [6, 26, 16, 18, 16, 22, 16, 16, 12, 40]
# The wrong tone to try for each tone: the one most often confused with it.
WRONG_FOR = {"mid": "low", "low": "mid", "falling": "high", "high": "falling", "rising": "mid"}


def tone_of(paiboon):
    marks = unicodedata.normalize("NFD", paiboon)
    if "̌" in marks:
        return "rising"
    if "̂" in marks:
        return "falling"
    if "́" in marks:
        return "high"
    if "̀" in marks:
        return "low"
    return "mid"


def trial_words():
    """Up to PER_TONE one-syllable words per tone, standard Thai topics first."""
    by_tone = {tone: [] for tone in TONE_ORDER}
    # Isan words last: the model voice speaks Central Thai.
    topics = sorted(app.VOCABULARY, key=lambda t: t == "isan_dialect")
    for topic in topics:
        for w in app.VOCABULARY[topic]:
            p = w["paiboon"]
            if " " in p or "-" in p:
                continue
            if not thai_audio.audio_exists(app.app.static_folder, w["thai"]):
                continue
            tone = tone_of(p)
            if len(by_tone[tone]) < PER_TONE:
                by_tone[tone].append((topic, w["thai"], p))
    return [(tone, *word) for tone in TONE_ORDER for word in by_tone[tone]]


def build(out_path):
    rows = trial_words()
    wb = Workbook()
    intro = wb.active
    intro.title = "อ่านก่อน · Read me"
    intro.column_dimensions["A"].width = 110
    intro.append(["ทดลอง 🎤 Say it · ThaiBridge AI · Pronunciation Analysis test"])
    intro["A1"].font = Font(name=THAI_FONT, size=16, bold=True, color="7A3B00")
    intro.append([])
    for para in INTRO_TH + [""] + INTRO_EN + ["", "{} words to try.".format(len(rows))]:
        intro.append([para])
    for row in intro.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.font = Font(name=THAI_FONT, size=12)

    ws = wb.create_sheet("ทดลอง · Trials")
    ws.append(HEADERS)
    for n, (tone, topic, thai, paiboon) in enumerate(rows, 1):
        wrong = WRONG_FOR[tone]
        ws.append([n, "/exercise/" + topic, thai, paiboon,
                   "{} ({})".format(TONE_TH[tone], tone),
                   "{} ({})".format(TONE_TH[wrong], wrong), "", "", "", ""])
    for i, w in enumerate(WIDTHS):
        ws.column_dimensions[chr(65 + i)].width = w
    head_fill = PatternFill("solid", fgColor="F6EFE2")
    for cell in ws[1]:
        cell.font = Font(name=THAI_FONT, bold=True)
        cell.fill = head_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.font = Font(name=THAI_FONT, size=11)
        row[2].font = Font(name=THAI_FONT, size=14)
    fair = DataValidation(type="list", formula1='"✓,✗,?"', allow_blank=True)
    ws.add_data_validation(fair)
    fair.add("I2:I{}".format(len(rows) + 1))
    ws.freeze_panes = "D2"

    wb.save(out_path)
    return rows


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--out", default=os.path.join("docs", "pronunciation-review-2026-10-27.xlsx"))
    args = parser.parse_args()
    out = args.out if os.path.isabs(args.out) else os.path.join(REPO_ROOT, args.out)
    rows = build(out)
    print("Wrote {}: {} words".format(args.out, len(rows)))
    for tone in TONE_ORDER:
        print("  {:8} {}".format(tone, ", ".join(r[2] for r in rows if r[0] == tone)))


if __name__ == "__main__":
    main()
