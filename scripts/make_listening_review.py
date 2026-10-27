"""Build the native speaker's Listening review spreadsheet.

Batch 1 checked the Thai TEXT. This sheet checks two new things:

  * the AUDIO — every clip the Listening Quiz and the Listen & answer
    conversations play. The voice is computer-generated (edge-tts), so a
    word can be written right and still be said wrong: a wrong tone, a wrong
    vowel length, a robotic run of words. Each row links to the clip on the
    live site, so the reviewer just clicks and listens.
  * the COMPREHENSION QUESTIONS in listening.py — are they fair, and is the
    right answer really what the Thai says?

Read from the live data, so the reviewer hears exactly what learners hear.

    python scripts/make_listening_review.py
    python scripts/make_listening_review.py --out docs/listening-review-2026-10-27.xlsx

The reviewer marks each row ✓ / ✗ / ? and writes a note for anything wrong;
those notes are what gets fixed (re-record, change the Thai, or reword a
question).
"""

import argparse
import os
import sys

# Importing app connects to the database at startup. Empty, not popped — the
# same reason as tests/conftest.py: dotenv would load the LIVE url back in.
os.environ["DATABASE_URL"] = ""

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, REPO_ROOT)

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

import app  # noqa: E402
import listening  # noqa: E402
import thai_audio  # noqa: E402

LIVE = "https://thaibridge-ai.smoald.com/static/"
THAI_FONT = "Leelawadee UI"

INTRO_TH = [
    "สวัสดีครับ 🙏",
    "ชุดนี้เป็นการตรวจ \"เสียง\" ในแบบฝึกฟังของแอป ThaiBridge AI ไม่ใช่ตรวจตัวหนังสือ",
    "เสียงทั้งหมดเป็นเสียงคอมพิวเตอร์ ไม่ใช่คนไทยอ่าน จึงอาจอ่านวรรณยุกต์ผิด สระสั้นยาวผิด "
    "หรือฟังดูไม่เป็นธรรมชาติ ขอให้ช่วยฟังแต่ละคำว่าออกเสียงถูกไหมครับ",
    "วิธีตรวจ: ในแผ่นงาน \"ฟังเสียง\" กดลิงก์ในช่อง ฟัง เพื่อฟังเสียง แล้วในช่อง ผล เลือก ✓ ถ้าถูก "
    "✗ ถ้าผิด หรือ ? ถ้าไม่แน่ใจ ถ้าผิด ช่วยเขียนว่าผิดตรงไหนในช่อง หมายเหตุ "
    "เช่น \"วรรณยุกต์ควรเป็นเสียงโท\" ครับ",
    "แผ่นงาน \"คำถาม\" เป็นคำถามภาษาอังกฤษหลังฟังบทสนทนา ช่วยดูว่าคำตอบที่ถูกตรงกับที่พูดในภาษาไทยไหมครับ",
    "ขอบพระคุณมากครับ",
]
INTRO_EN = [
    "This review is about how the listening exercises SOUND, not how the Thai is written "
    "(that was Batch 1).",
    "Every clip is a computer voice, not a Thai speaker, so it can say a correctly written "
    "word wrongly: a wrong tone, a wrong vowel length, or an unnatural run of words.",
    "How: on the Audio sheet, click the link under Listen, then choose ✓ if it sounds right, "
    "✗ if it is wrong, or ? if unsure. For anything wrong, write what is wrong in the Note "
    "column, e.g. \"should be a falling tone\". Skip anything you like.",
    "The Questions sheet has the English comprehension questions asked after each "
    "conversation. Please check the right answer really matches what the Thai says, and "
    "that no other choice could also be right.",
]

AUDIO_HEADERS = ["#", "หน้า · Page", "ภาษาไทย · Thai", "ThaiBridge Spelling",
                 "ความหมาย · Meaning", "ฟัง · Listen", "ผล · Result",
                 "หมายเหตุ · Note (what is wrong)"]
AUDIO_WIDTHS = [6, 30, 36, 30, 34, 12, 10, 40]

QUESTION_HEADERS = ["#", "บทสนทนา · Conversation", "คำถาม · Question",
                    "คำตอบที่ถูก · Right answer", "ตัวเลือกอื่น · Other choices",
                    "ผล · Result", "หมายเหตุ · Note"]
QUESTION_WIDTHS = [6, 30, 40, 32, 46, 10, 40]


def audio_rows():
    """Every clip a listening exercise plays, each Thai text once."""
    rows, seen = [], set()

    def add(page, line):
        thai = line["thai"]
        if thai in seen or not thai_audio.audio_exists(app.app.static_folder, thai):
            return
        seen.add(thai)
        rows.append((page, thai, line.get("paiboon", ""), line.get("english", ""),
                     LIVE + thai_audio.audio_static_path(thai)))

    for conv in app.CONVERSATIONS.values():
        for line in conv["dialogue"]:
            add("Sentences › " + conv["title"], line)
    for topic, words in app.VOCABULARY.items():
        for word in words:
            add("Listening Quiz › " + topic.replace("_", " ").title(), word)
    return rows


def question_rows():
    rows = []
    for conv_id, questions in listening.QUESTIONS.items():
        title = app.CONVERSATIONS[conv_id]["title"]
        for q in questions:
            rows.append((title, q["question"], q["answer"], " · ".join(q["wrong"])))
    return rows


def style_sheet(ws, headers, widths, thai_col=None):
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    head_fill = PatternFill("solid", fgColor="F6EFE2")
    for cell in ws[1]:
        cell.font = Font(name=THAI_FONT, bold=True)
        cell.fill = head_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.font = Font(name=THAI_FONT, size=11)
        if thai_col is not None:
            row[thai_col].font = Font(name=THAI_FONT, size=14)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = "A1:{}{}".format(chr(64 + len(headers)), ws.max_row)


def add_result_dropdown(ws, column, last_row):
    result = DataValidation(type="list", formula1='"✓,✗,?"', allow_blank=True)
    ws.add_data_validation(result)
    result.add("{0}2:{0}{1}".format(column, last_row))


def build(out_path):
    audio = audio_rows()
    questions = question_rows()

    wb = Workbook()
    intro = wb.active
    intro.title = "อ่านก่อน · Read me"
    intro.column_dimensions["A"].width = 110
    intro.append(["ตรวจเสียงแบบฝึกฟัง · ThaiBridge AI · Listening review"])
    intro["A1"].font = Font(name=THAI_FONT, size=16, bold=True, color="7A3B00")
    intro.append([])
    for para in INTRO_TH + [""] + INTRO_EN + [""]:
        intro.append([para])
    intro.append(["{} clips to listen to, and {} questions to check.".format(
        len(audio), len(questions))])
    for row in intro.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.font = Font(name=THAI_FONT, size=12)

    ws = wb.create_sheet("ฟังเสียง · Audio")
    ws.append(AUDIO_HEADERS)
    for n, (page, thai, spelling, meaning, url) in enumerate(audio, 1):
        ws.append([n, page, thai, spelling, meaning, "▶ ฟัง · Listen", "", ""])
        link = ws.cell(row=n + 1, column=6)
        link.hyperlink = url
    style_sheet(ws, AUDIO_HEADERS, AUDIO_WIDTHS, thai_col=2)
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=6):
        row[0].font = Font(name=THAI_FONT, size=11, color="0563C1", underline="single")
    add_result_dropdown(ws, "G", len(audio) + 1)

    qs = wb.create_sheet("คำถาม · Questions")
    qs.append(QUESTION_HEADERS)
    for n, row in enumerate(questions, 1):
        qs.append([n, *row, "", ""])
    style_sheet(qs, QUESTION_HEADERS, QUESTION_WIDTHS)
    add_result_dropdown(qs, "F", len(questions) + 1)

    wb.save(out_path)
    return audio, questions


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--out", default=os.path.join("docs", "listening-review-2026-10-27.xlsx"))
    args = parser.parse_args()
    out = args.out if os.path.isabs(args.out) else os.path.join(REPO_ROOT, args.out)
    audio, questions = build(out)
    print("Wrote {}: {} clips, {} questions".format(args.out, len(audio), len(questions)))


if __name__ == "__main__":
    main()
