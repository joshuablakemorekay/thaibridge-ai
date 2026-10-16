"""Build the whole-app Thai review workbook for a native-speaker reviewer.

One workbook, four tabs:
  * How to use      — short instructions in English and Thai;
  * Checklist       — every page and part of a page that carries Thai, with a
                      link, so the reviewer can tick each one off and nothing
                      is skipped by accident;
  * Findings        — blank rows for anything they spot while browsing;
  * All Thai lines  — every Thai line a learner can meet, page by page, read
                      from the live data (not a copy that may have drifted).

It builds on make_batch1_review.py (the core learning pages) and adds what
that batch left for later: Monk Mode, the Tour Guide trips, the Dhamma talks,
the journal and the Thai typed into those pages' HTML.

Left out on purpose:
  * the chanting book — a separate, much bigger job;
  * Pra Kru Bob's writings — his own words, published unaltered with consent;
  * the Yaitron dictionary entries — a published third-party dictionary;
  * the female form of a line where it differs only in ค่ะ/คะ/ดิฉัน (the
    pronoun question is asked once instead of doubling the list).

    python scripts/make_full_review.py --out path/to/review.xlsx
"""

import argparse
import glob
import importlib.util
import os
import re
import sys

os.environ["DATABASE_URL"] = ""  # importing app must never reach the live database

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, REPO_ROOT)
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

_spec = importlib.util.spec_from_file_location(
    "make_batch1_review", os.path.join(REPO_ROOT, "scripts", "make_batch1_review.py"))
batch1 = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(batch1)

app = batch1.app
import thai_reading  # noqa: E402
import tour_trips  # noqa: E402

SITE = "https://thaibridge-ai.smoald.com"
THAI_FONT = "Leelawadee UI"
HEAD_FILL = PatternFill("solid", fgColor="7A3E00")
HEAD_FONT = Font(name=THAI_FONT, bold=True, color="FFFFFF")
JOSH_FILL = PatternFill("solid", fgColor="EDEDED")
WRAP = Alignment(wrap_text=True, vertical="top")

# Thai typed into a page's HTML: only the chanting book, the print edition and
# Pra Kru Bob's essays stay out (see the docstring for why).
batch1.TEMPLATE_SKIP = re.compile(r"(chant|print[\\/]|bob_)")


def en(title):
    """Titles are sometimes a plain string, sometimes {'en': ..., 'th': ...}."""
    return title.get("en", "") if isinstance(title, dict) else (title or "")


# ---------------------------------------------------------------- checklist

def checklist():
    """(section, part, address, what to check) for every page with Thai on it."""
    rows = []

    def add(section, part, url, note=""):
        rows.append((section, part, SITE + url, note))

    add("Alphabet", "All 44 consonants, their names and the quiz", "/alphabet")
    add("Survival Thai", "Every set: " + ", ".join(s["title"] for s in batch1.survival.SETS), "/survival",
        "Try the formal / neutral / casual switch too")
    add("Vocabulary", "Topic list", "/learn")
    for key, words in app.VOCABULARY.items():
        add("Vocabulary", "{} ({} words)".format(key.replace("_", " ").title(), len(words)),
            "/exercise/" + key, "Words and the quiz")
    add("Lessons", "Lesson list", "/lessons")
    for lesson in app.LESSONS:
        add("Lessons", "Lesson {}: {}".format(lesson["id"], en(lesson["title"])),
            "/lesson/{}".format(lesson["id"]))
    for part in ("Sentence patterns", "Conversations", "Conversational phrases",
                 "Build the sentence (every level)", "How Thai is really said (note box)"):
        add("Sentences & Conversations", part, "/sentences")
    add("Grammar Guide", "Whole page", "/grammar")
    add("Formality Levels", "Whole page", "/formality")
    add("Formality Levels", "Registers guide", "/register")
    add("Gender Guide", "Whole page", "/gender-examples")
    add("Tones & Consonant Classes", "Whole page, including the tone-deaf help", "/tones-classes")
    add("Vowels & Syllables", "Whole page", "/vowels_syllables")
    add("Read & Write Script", "Vowel positions and Dhamma words", "/read_write")
    for story in thai_reading.STORIES:
        add("Read & Write Script", "Story: " + story.get("title_en", story.get("key", "")), "/read_write")
    add("ThaiBridge Spelling", "Whole page", "/paiboon", "The tone and sound rules")
    add("Dictionary", "Reference part only", "/dictionary",
        "The search results come from a published dictionary (Yaitron), no need to check those")
    add("Thai Culture", "Every topic and story", "/culture")
    add("Theravada Dhamma", "Whole page", "/theravada")
    for talk in app.DHAMMA_TALKS:
        add("Theravada Dhamma", "Dhamma talk: " + en(talk.get("title")), "/dhamma-talk/" + talk["slug"])
    add("Theravada Dhamma", "Dhamma and culture", "/dhamma-and-culture")
    add("Theravada Dhamma", "Practising anywhere", "/practising-anywhere")
    add("Meditation", "Whole page", "/meditation")
    add("Tour Guide Thai", "Vocabulary", "/tour-guide")
    for trip in tour_trips.TRIPS:
        add("Tour Guide Thai", "Trip: " + trip.get("title", trip["key"]), "/tour-guide",
            "Play the trip through: what you hear and the answers")
    add("Business Thai", "Whole page", "/business-thai")
    add("Monk Mode", "Lesson list", "/monk/lessons", "Needs the Monk Mode code first")
    for topic in app.MONK_TOPICS:
        add("Monk Mode", "Lesson: " + en(topic.get("title")), "/monk/lesson/" + topic["topic"])
    add("Monk Mode", "Pronunciation guide", "/monk/pronunciation")
    add("AI Tutor", "Ask it a few questions", "/chat",
        "Optional. Its answers are written fresh each time, so tell me if its Thai sounds wrong")
    for part, url in (("Home page", "/"), ("About", "/about"), ("How to use the app", "/instructions"),
                      ("Plans and prices", "/premium"), ("Progress", "/progress"), ("Journal", "/journal")):
        add("Other pages", part, url, "Only the Thai on these pages")
    return rows


# ---------------------------------------------------------------- all lines

def page_of_template():
    """Map each template file to the address that shows it, from app.py itself."""
    src = open(os.path.join(REPO_ROOT, "app.py"), encoding="utf-8").read()
    found = {}
    for block in re.split(r"\n@app\.route\(", src)[1:]:
        url = re.match(r"'([^']*)'", block)
        tpl = re.search(r"render_template\(\s*'([^']+)'", block.split("\n\n\n")[0])
        if url and tpl and tpl.group(1) not in found:
            found[tpl.group(1)] = url.group(1)
    return found


ATTR_THAI = re.compile(r"""(?:title|aria-label|placeholder|alt)="([^"]*[฀-๿][^"]*)"|"""
                       r"""\{%\s*block title\s*%\}(.*?[฀-๿].*?)\{%""")


def attribute_rows(sheet):
    """Thai in tooltips, page titles and labels, which the text pass strips out."""
    for path in sorted(glob.glob(os.path.join(REPO_ROOT, "templates", "**", "*.html"), recursive=True)):
        rel = os.path.relpath(path, REPO_ROOT).replace("\\", "/")
        if batch1.TEMPLATE_SKIP.search(rel):
            continue
        for line in open(path, encoding="utf-8"):
            for attr, title in ATTR_THAI.findall(line):
                sheet.add("Written into the page", rel, "tooltip or page title", attr or title,
                          "", "(note: check the Thai inside)")


def all_lines():
    sheet = batch1.Sheet()
    for page, url, data in batch1.PAGES:
        sheet.walk(page, url, data)
        if page == "Sentences: conversational phrases":
            batch1.builder_rows(sheet)
        if page == "Register switcher (shown on several pages)":
            batch1.register_switch_rows(sheet)
        if page == "Vowels & syllables":
            batch1.consonant_rows(sheet)
    sheet.walk("Dictionary: reference", "/dictionary", app.DICTIONARY_REFERENCE)
    for topic in app.MONK_TOPICS:
        sheet.walk("Monk Mode: " + en(topic.get("title")), "/monk/lesson/" + topic["topic"], topic)
    sheet.walk("Tour Guide: particles", "/tour-guide", tour_trips.PARTICLES)
    for trip in tour_trips.TRIPS:
        sheet.walk("Tour Guide trip: " + trip.get("title", trip["key"]), "/tour-guide", trip)
    for talk in app.DHAMMA_TALKS:
        sheet.walk("Dhamma talk: " + en(talk.get("title")), "/dhamma-talk/" + talk["slug"], talk)
    sheet.walk("Journal", "/journal", app.JOURNAL_ENTRIES)
    # The batch-1 pass only lists switcher versions that have a 'thai' field;
    # notes-only versions (e.g. the monastic tips) carry Thai too.
    sheet.walk("Register switcher (shown on several pages)", "/survival",
               batch1.thai_registers.VARIANTS)
    for c in batch1.thai_consonants.CONSONANTS:
        spoken = c.get("speak")
        if spoken and spoken != c["name_thai"]:
            sheet.add("Alphabet: consonants", "/alphabet", "what the recording says",
                      spoken, "", c["name"])
    # Quiz feedback built in Python, not stored in any data table.
    for text in ("ถูกต้อง! (Correct!)", "ไม่ถูกต้อง. Answer: …"):
        sheet.add("Quiz feedback (every quiz)", "", "after an answer", text)

    before = len(sheet.rows)
    batch1.template_rows(sheet)
    attribute_rows(sheet)
    pages = page_of_template()
    # Scripts and partials that page_of_template() can't see being rendered.
    pages.update({"read_write.js": "/read_write", "tour_trips.js": "/tour-guide",
                  "alphabet.js": "/alphabet", "exercise.html": "/learn",
                  "vowel_practice.html": "/vowels_syllables",
                  "how_thai_is_said.html": "/sentences",
                  "monk_accent_switch.html": "/monk/lessons"})
    rows = sheet.rows[:before]
    for page, rel, where, thai, spelling, meaning in sheet.rows[before:]:
        name = os.path.basename(rel)
        url = pages.get(name)
        if url is None:
            page, url = "Shared by several pages ({})".format(name), ""
        else:
            page = "Written into the page"
        rows.append((page, url, where, thai, spelling, meaning))
    return [(p, SITE + u if u.startswith("/") else u, w, t, s, m) for p, u, w, t, s, m in rows]


# ---------------------------------------------------------------- workbook

FINDING_PAGES = [
    "Alphabet", "Survival Thai", "Vocabulary", "Lessons", "Sentences & Conversations",
    "Grammar Guide", "Formality Levels", "Gender Guide", "Tones & Consonant Classes",
    "Vowels & Syllables", "Read & Write Script", "ThaiBridge Spelling", "Dictionary",
    "Thai Culture", "Theravada Dhamma", "Meditation", "Tour Guide Thai", "Business Thai",
    "Monk Mode", "AI Tutor", "Other pages",
]
TYPES = ["1. Wrong / ผิด", "2. Unnatural / ไม่เป็นธรรมชาติ", "3. Spelling / คำอ่านผิด",
         "4. Culturally off / ไม่เหมาะกับวัฒนธรรม"]
STATUS = ["To do", "Fixed", "Question for the reviewer", "Not changing"]

HOW_TO = [
    ("ThaiBridge AI: Thai review", "title"),
    ("", ""),
    ("English", "bold"),
    ("Thank you for checking the Thai! There are three tabs to use:", ""),
    ("• Checklist: every page in the app, with a link. Tick each one when you have checked it, "
     "so we both know nothing has been missed.", ""),
    ("• Findings: whenever you find something while using the app, add one row here "
     "(page, type, what it says now, what it should say, why).", ""),
    ("• All Thai lines: every Thai line in the app, page by page. Optional. Use it if you would "
     "rather read a list, or to double-check. Mark a line ✓ if it is fine, ✗ and write the "
     "correction if not, or ? if unsure.", ""),
    ("You don't need to fill in the grey columns. Those are for me, so you can see what I've fixed.", ""),
    ("Lines are mostly shown as a man says them (ครับ / ผม). The app changes them to ค่ะ / คะ for women.", ""),
    ("Not included: the Chanting Book (a separate job) and Pra Kru Bob's writings (his own words).", ""),
    ("Skip anything you're unsure about. No deadline, go at your own pace.", ""),
    ("", ""),
    ("ภาษาไทย", "bold"),
    ("ขอบคุณมากที่ช่วยตรวจภาษาไทยครับ ในไฟล์นี้มีสามแท็บที่ใช้ครับ", ""),
    ("• Checklist: รายชื่อทุกหน้าในแอปพร้อมลิงก์ ตรวจหน้าไหนเสร็จแล้วติ๊กไว้ จะได้รู้ว่าไม่มีหน้าไหนตกหล่นครับ", ""),
    ("• Findings: ระหว่างใช้แอป เจอจุดไหนเพิ่มหนึ่งแถวได้เลยครับ (หน้า ประเภท ตอนนี้เขียนว่า ควรแก้เป็น เหตุผล)", ""),
    ("• All Thai lines: ภาษาไทยทุกบรรทัดในแอป เรียงตามหน้า ไม่บังคับครับ ใช้ถ้าสะดวกอ่านเป็นรายการ "
     "หรือใช้ตรวจซ้ำ ถูกแล้วเลือก ✓ ต้องแก้เลือก ✗ แล้วเขียนแบบที่ถูก ไม่แน่ใจเลือก ?", ""),
    ("ไม่ต้องกรอกช่องสีเทานะครับ ช่องนั้นผมจะกรอกเอง เพื่อให้เห็นว่าแก้อะไรไปแล้วบ้าง", ""),
    ("ประโยคส่วนใหญ่แสดงแบบผู้ชาย (ครับ / ผม) แอปจะเปลี่ยนเป็น ค่ะ / คะ ให้ผู้หญิงเองครับ", ""),
    ("ไม่รวม: หนังสือสวดมนต์ (เป็นงานแยก) และบทความของพระครูบ็อบ (เป็นถ้อยคำของท่านเอง)", ""),
    ("ข้อไหนไม่แน่ใจ ข้ามได้เลยครับ ไม่มีกำหนดส่ง ทำตามสะดวกครับ", ""),
]


def header(ws, cols):
    for i, (text, width) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=i, value=text)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[c.column_letter].width = width
    ws.row_dimensions[1].height = 36


def dropdown(ws, values, rng):
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
    dv.showErrorMessage = False
    ws.add_data_validation(dv)
    dv.add(rng)


def build(out):
    wb = Workbook()

    info = wb.active
    info.title = "How to use - วิธีใช้"
    info.column_dimensions["A"].width = 110
    for i, (text, style) in enumerate(HOW_TO, start=1):
        c = info.cell(row=i, column=1, value=text)
        c.alignment = WRAP
        c.font = Font(name=THAI_FONT, size=14 if style == "title" else 11, bold=bool(style))

    checks = checklist()
    ws = wb.create_sheet("Checklist")
    header(ws, [("#", 5), ("Section\nหมวด", 24), ("Page or part\nหน้า / ส่วน", 44),
                ("Link\nลิงก์", 46), ("Tip\nหมายเหตุ", 40), ("Checked?\nตรวจแล้ว?", 14),
                ("Notes\nหมายเหตุ", 30)])
    for n, (section, part, url, note) in enumerate(checks, start=1):
        ws.append([n, section, part, url, note, "", ""])
        ws.cell(row=n + 1, column=4).hyperlink = url
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
    dropdown(ws, ["✓ Done / เสร็จ", "Part done / บางส่วน"], "F2:F{}".format(len(checks) + 1))
    ws.freeze_panes = "C2"

    ws = wb.create_sheet("Findings")
    header(ws, [("#", 5), ("Page\nหน้า", 22), ("Type\nประเภท", 26),
                ("What it says now\nตอนนี้เขียนว่า", 34), ("What it should say\nควรแก้เป็น", 34),
                ("Why (optional)\nเหตุผล (ถ้ามี)", 34), ("Screenshot on WhatsApp?\nส่งภาพทาง WhatsApp?", 16),
                ("Status (Josh)\nสถานะ", 18), ("Josh's notes\nหมายเหตุจาก Josh", 30)])
    last = 400
    for r in range(2, last + 1):
        ws.cell(row=r, column=1, value=r - 1)
        for col in range(2, 10):
            ws.cell(row=r, column=col).alignment = WRAP
        for col in (8, 9):
            ws.cell(row=r, column=col).fill = JOSH_FILL
    dropdown(ws, FINDING_PAGES, "B2:B{}".format(last))
    dropdown(ws, TYPES, "C2:C{}".format(last))
    dropdown(ws, ["Yes / ใช่", "No / ไม่"], "G2:G{}".format(last))
    dropdown(ws, STATUS, "H2:H{}".format(last))
    ws.freeze_panes = "B2"

    lines = all_lines()
    ws = wb.create_sheet("All Thai lines")
    header(ws, [("#", 6), ("Page\nหน้า", 26), ("Link\nลิงก์", 30), ("Where on the page\nส่วน", 28),
                ("Thai\nภาษาไทย", 44), ("ThaiBridge Spelling", 30), ("Meaning\nความหมาย", 34),
                ("OK?\nผล", 8), ("Correction\nแก้เป็น", 36), ("Comment\nหมายเหตุ", 28),
                ("Status (Josh)\nสถานะ", 16)])
    for n, row in enumerate(lines, start=1):
        ws.append([n, *row, "", "", "", ""])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
        row[4].font = Font(name=THAI_FONT, size=13)
        row[5].font = Font(italic=True, color="8A5A00")
        row[10].fill = JOSH_FILL
    end = len(lines) + 1
    dropdown(ws, ["✓", "✗", "?"], "H2:H{}".format(end))
    dropdown(ws, STATUS, "K2:K{}".format(end))
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = "A1:K{}".format(end)

    wb.save(out)
    return checks, lines


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--out", required=True)
    args = parser.parse_args()
    checks, lines = build(args.out)
    pages = {}
    for r in lines:
        pages[r[0]] = pages.get(r[0], 0) + 1
    for page, n in pages.items():
        print("{:5}  {}".format(n, page))
    print("Checklist: {} rows. All Thai lines: {} rows. Wrote {}".format(len(checks), len(lines), args.out))


if __name__ == "__main__":
    main()
